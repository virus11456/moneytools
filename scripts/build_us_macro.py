#!/usr/bin/env python3
"""Build public US macro data from auditable public sources.

FRED remains the primary source.  CFTC data and market-price-derived series are
added only when their exact source and formula can be disclosed on the site.
Unavailable licensed surveys remain absent instead of being approximated.
"""
from __future__ import annotations
import csv, html, io, json, math, re, subprocess, zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "public/macro-data/us-macro.json"
SERIES = {
 "ICSA":("level",.001),"CCSA":("level",.001),"PAYEMS":("change",1),
 "PCE":("yoy",1),"PI":("yoy",1),"PSAVERT":("level",1),
 "SPCS20RSA":("level",1),"MSPNHSUS":("level",.001),"HSN1F":("yoy",1),"EXHOSLUSM495S":("level",1),
 "TOTALSA":("level",1),"RSAFS":("yoy",1),"MRTSSM4541USN":("yoy",1),
 "HOUST":("yoy",1),"PERMIT":("yoy",1),"HOSSUPUSM673N":("level",1),
 "DGORDER":("yoy",1),"NEWORDER":("yoy",1),"WEI":("level",1),"GDP":("yoy4",1),
 "PPIFIS":("yoy",1),"CPIAUCSL":("yoy",1),"CPILFESL":("yoy",1),
 "FEDFUNDS":("level",1),"DGS2":("level",1),"DGS5":("level",1),"DGS10":("level",1),"SOFR90DAYAVG":("level",1),
 "BAMLH0A3HYC":("level",1),"VIXCLS":("level",1),"SP500":("level",1),
}

MARKET_SYMBOLS = {
 "COPPER": "HG=F", "GOLD": "GC=F", "OIL": "CL=F", "NDX": "^NDX",
 "VTI": "VTI", "VOX": "VOX", "VCR": "VCR", "VDC": "VDC",
 "VDE": "VDE", "VFH": "VFH", "VHT": "VHT", "VIS": "VIS",
 "VAW": "VAW", "VNQ": "VNQ", "VGT": "VGT", "VPU": "VPU",
}
SECTOR_ETFS = ("VOX", "VCR", "VDC", "VDE", "VFH", "VHT", "VIS", "VAW", "VNQ", "VGT", "VPU")
CFTC_CODES = {"SP500_COT": "13874+", "NASDAQ_COT": "20974+"}

def fetch(series: str):
    result=subprocess.run(["curl","--fail","--location","--silent","--show-error","--max-time","30",f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={series}"],check=True,capture_output=True,text=True)
    raw=result.stdout
    rows=[]
    for row in csv.DictReader(io.StringIO(raw)):
        try: v=float(row.get(series,""))
        except (ValueError,TypeError): continue
        if math.isfinite(v): rows.append((row.get("observation_date") or row.get("DATE"),v))
    return rows

def fetch_bytes(url: str, timeout: int = 90, headers: dict[str, str] | None = None) -> bytes:
    command = [
        "curl", "--fail", "--location", "--silent", "--show-error", "--compressed",
        "--retry", "3", "--retry-all-errors", "--retry-delay", "1", "--max-time", str(timeout),
    ]
    for name, value in (headers or {}).items():
        command.extend(["--header", f"{name}: {value}"])
    command.append(url)
    result = subprocess.run(
        command,
        check=True, capture_output=True,
    )
    return result.stdout

def transform(rows, mode, scale):
    lag=4 if mode=="yoy4" else 12
    out=[]
    for i,(date,value) in enumerate(rows):
        if mode=="change":
            if i<1: continue
            value=value-rows[i-1][1]
        elif mode in ("yoy","yoy4"):
            if i<lag or rows[i-lag][1]==0: continue
            value=(value/rows[i-lag][1]-1)*100
        out.append({"date":date,"value":round(value*scale,4)})
    # Keep twenty-five years and thin daily series to Friday/last observation per week.
    cutoff="2000-01-01"; out=[p for p in out if p["date"]>=cutoff]
    if len(out)>2500:
        weekly={}
        for p in out:
            d=datetime.fromisoformat(p["date"]); weekly[f"{d.isocalendar().year}-{d.isocalendar().week:02}"]=p
        out=list(weekly.values())
    return out

def points_result(points):
    return {"latestDate": points[-1]["date"] if points else None, "points": points}

def ratio_points(numerator, denominator):
    """Return a same-day price ratio without forward-filling either market."""
    den = {date: value for date, value in denominator if value and math.isfinite(value)}
    points = [
        {"date": date, "value": round(value / den[date], 6)}
        for date, value in numerator
        if date in den and den[date] and math.isfinite(value)
    ]
    return thin_points(points)

def thin_points(points):
    points = [p for p in points if p["date"] >= "2000-01-01" and math.isfinite(p["value"])]
    if len(points) <= 2500:
        return points
    weekly = {}
    for point in points:
        date = datetime.fromisoformat(point["date"])
        weekly[f"{date.isocalendar().year}-{date.isocalendar().week:02}"] = point
    return list(weekly.values())

def fetch_market_series():
    """Download adjusted closes once, then derive all transparent market ratios."""
    import yfinance as yf

    requested = list(MARKET_SYMBOLS.values())
    frame = yf.download(
        requested, start="2000-01-01", auto_adjust=True, progress=False,
        group_by="ticker", threads=True,
    )
    prices = {}
    for key, symbol in MARKET_SYMBOLS.items():
        try:
            close = frame[symbol]["Close"] if len(requested) > 1 else frame["Close"]
            rows = []
            for index, value in close.dropna().items():
                number = float(value)
                if math.isfinite(number):
                    rows.append((index.strftime("%Y-%m-%d"), number))
            prices[key] = rows
        except Exception:
            prices[key] = []

    output = {
        "COPPER_GOLD": points_result(ratio_points(prices["COPPER"], prices["GOLD"])),
        "OIL_GOLD": points_result(ratio_points(prices["OIL"], prices["GOLD"])),
        "NDX": points_result(thin_points([{"date": date, "value": round(value, 4)} for date, value in prices["NDX"]])),
    }
    for symbol in SECTOR_ETFS:
        output[f"{symbol}_VTI"] = points_result(ratio_points(prices[symbol], prices["VTI"]))
    return output

def fetch_breadth_series():
    """Calculate current-constituent S&P 500 breadth with an explicit bias note in UI."""
    import yfinance as yf

    raw = fetch_bytes("https://raw.githubusercontent.com/datasets/s-and-p-500-companies/main/data/constituents.csv").decode("utf-8-sig")
    tickers = []
    for row in csv.DictReader(io.StringIO(raw)):
        symbol = (row.get("Symbol") or "").strip().replace(".", "-")
        if symbol:
            tickers.append(symbol)
    if len(tickers) < 450:
        raise ValueError(f"S&P 500 constituent list unexpectedly short: {len(tickers)}")
    frame = yf.download(
        tickers, start="2015-01-01", auto_adjust=True, progress=False,
        group_by="column", threads=True,
    )
    close = frame["Close"]
    output = {}
    for window, series_id in ((50, "SPX_ABOVE_50"), (200, "SPX_ABOVE_200")):
        average = close.rolling(window=window, min_periods=window).mean()
        eligible = close.notna() & average.notna()
        denominator = eligible.sum(axis=1)
        ratio = ((close > average) & eligible).sum(axis=1).div(denominator.where(denominator > 0)).mul(100).dropna()
        points = [
            {"date": index.strftime("%Y-%m-%d"), "value": round(float(value), 4)}
            for index, value in ratio.items() if math.isfinite(float(value))
        ]
        output[series_id] = points_result(thin_points(points))
    return output

def parse_cftc_zip(payload: bytes):
    """Parse TFF futures-only leveraged-money net positions by contract code."""
    output = {series_id: {} for series_id in CFTC_CODES}
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        filename = archive.namelist()[0]
        with archive.open(filename) as raw:
            rows = csv.DictReader(io.TextIOWrapper(raw, encoding="utf-8-sig", errors="replace"))
            for row in rows:
                code = (row.get("CFTC_Contract_Market_Code") or "").strip()
                for series_id, wanted_code in CFTC_CODES.items():
                    if code != wanted_code:
                        continue
                    try:
                        long_value = float(row["Lev_Money_Positions_Long_All"])
                        short_value = float(row["Lev_Money_Positions_Short_All"])
                    except (KeyError, TypeError, ValueError):
                        continue
                    raw_date = row.get("Report_Date_as_YYYY-MM-DD", "").strip()
                    date = None
                    for pattern in ("%Y-%m-%d", "%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y"):
                        try:
                            date = datetime.strptime(raw_date, pattern).strftime("%Y-%m-%d")
                            break
                        except ValueError:
                            pass
                    if date:
                        output[series_id][date] = long_value - short_value
    return output

def fetch_cftc_series():
    combined = {series_id: {} for series_id in CFTC_CODES}
    urls = ["https://www.cftc.gov/files/dea/history/fin_fut_txt_2006_2016.zip"]
    urls.extend(
        f"https://www.cftc.gov/files/dea/history/fut_fin_txt_{year}.zip"
        for year in range(2017, datetime.now(timezone.utc).year + 1)
    )
    errors = []
    for url in urls:
        try:
            parsed = parse_cftc_zip(fetch_bytes(url))
            for series_id, values in parsed.items():
                combined[series_id].update(values)
        except Exception as exc:
            errors.append(f"{url.rsplit('/', 1)[-1]}: {type(exc).__name__}")
    output = {}
    for series_id, values in combined.items():
        points = [{"date": date, "value": round(value, 4)} for date, value in sorted(values.items())]
        output[series_id] = points_result(points)
        if not points and errors:
            output[series_id]["error"] = "; ".join(errors)
    return output

def moving_average(points, window):
    values = []
    output = []
    for point in points:
        values.append(point["value"])
        if len(values) > window:
            values.pop(0)
        if len(values) == window:
            output.append({"date": point["date"], "value": round(sum(values) / window, 6)})
    return output

def moving_average_with_gap_reset(points, window, max_gap_days=10):
    """Moving average that never bridges a long publication gap."""
    values = []
    output = []
    previous = None
    for point in points:
        current = date.fromisoformat(point["date"])
        if previous and (current - previous).days > max_gap_days:
            values = []
        previous = current
        values.append(point["value"])
        if len(values) > window:
            values.pop(0)
        if len(values) == window:
            output.append({"date": point["date"], "value": round(sum(values) / window, 6)})
    return output

def fetch_aaii_series():
    """Load AAII's complete public historical workbook and derive its 20-week mean."""
    import xlrd

    payload = fetch_bytes(
        "https://www.aaii.com/files/surveys/sentiment.xls",
        headers={
            "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/140.0.0.0 Safari/537.36",
            "Referer": "https://www.aaii.com/sentimentsurvey/sent_results?reload=true",
        },
    )
    book = xlrd.open_workbook(file_contents=payload)
    sheet = book.sheet_by_name("SENTIMENT")
    points = []
    for row in range(5, sheet.nrows):
        try:
            raw_date = sheet.cell_value(row, 0)
            parts = xlrd.xldate_as_tuple(raw_date, book.datemode)
            reported = date(parts[0], parts[1], parts[2]).isoformat()
            bullish = float(sheet.cell_value(row, 1))
            bearish = float(sheet.cell_value(row, 3))
            spread = (bullish - bearish) * 100
        except (TypeError, ValueError, xlrd.XLDateError):
            continue
        if reported >= "2000-01-01" and math.isfinite(spread):
            points.append({"date": reported, "value": round(spread, 4)})
    points.sort(key=lambda item: item["date"])
    return {
        "AAII_SPREAD": points_result(points),
        "AAII_MA20": points_result(moving_average(points, 20)),
    }

def fetch_manheim_series():
    """Discover and parse Cox Automotive's latest official monthly MUVVI workbook."""
    from openpyxl import load_workbook

    query = "https://www.coxautoinc.com/wp-json/wp/v2/search?search=Manheim%20Used%20Vehicle%20Value%20Index&per_page=30"
    results = json.loads(fetch_bytes(query, headers={"User-Agent": "Stocktools/1.0"}))
    candidates = [
        item["url"] for item in results
        if re.fullmatch(r"Manheim Used Vehicle Value Index: (?!Mid-).+ Trends", item.get("title", ""), re.I)
    ]
    if not candidates:
        raise ValueError("latest Manheim monthly report not found")
    report = html.unescape(fetch_bytes(candidates[0], headers={"User-Agent": "Stocktools/1.0"}).decode("utf-8", "replace"))
    match = re.search(r'https://www\.coxautoinc\.com/wp-content/uploads/[^"\']+\.xlsx', report, re.I)
    if not match:
        raise ValueError("Manheim workbook link not found")
    workbook = load_workbook(io.BytesIO(fetch_bytes(match.group(0))), read_only=True, data_only=True)
    sheet = workbook["DATA"]
    points = []
    for raw_date, value, *_ in sheet.iter_rows(min_row=2, values_only=True):
        if not isinstance(raw_date, (date, datetime)):
            continue
        try:
            number = float(value)
        except (TypeError, ValueError):
            continue
        reported = raw_date.strftime("%Y-%m-%d")
        if reported >= "2000-01-01" and math.isfinite(number):
            points.append({"date": reported, "value": round(number, 4)})
    return {"MANHEIM": points_result(points)}

def parse_cboe_daily(payload: bytes, requested: str):
    """Extract the official equity put/call ratio from a Cboe daily RSC response."""
    body = payload.decode("utf-8", "replace")
    match = re.search(r'"name":"EQUITY PUT/CALL RATIO","value":"([0-9.]+)"', body)
    if not match:
        match = re.search(r'\\"name\\":\\"EQUITY PUT/CALL RATIO\\",\\"value\\":\\"([0-9.]+)\\"', body)
    if not match:
        return None
    return {"date": requested, "value": float(match.group(1))}

def fetch_cboe_daily(date_string: str):
    payload = fetch_bytes(
        f"https://www.cboe.com/us/options/market_statistics/daily/?dt={date_string}",
        timeout=45, headers={"RSC": "1", "User-Agent": "Stocktools/1.0"},
    )
    return parse_cboe_daily(payload, date_string)

def fetch_cboe_put_call():
    """Join Cboe's archive to its public daily pages and increment from the last build."""
    raw = fetch_bytes("https://cdn.cboe.com/resources/options/volume_and_call_put_ratios/equitypc.csv").decode("utf-8-sig", "replace")
    header = raw.find("DATE,CALL,PUT,TOTAL,P/C Ratio")
    if header < 0:
        raise ValueError("Cboe equity put/call header not found")
    points = []
    for row in csv.DictReader(io.StringIO(raw[header:])):
        try:
            reported = datetime.strptime(row["DATE"].strip(), "%m/%d/%Y").strftime("%Y-%m-%d")
            value = float(row["P/C Ratio"])
        except (KeyError, TypeError, ValueError):
            continue
        if math.isfinite(value):
            points.append({"date": reported, "value": value})
    by_date = {point["date"]: point for point in points}
    try:
        previous = json.loads(OUT.read_text(encoding="utf-8")).get("series", {}).get("CBOE_PC", {}).get("points", [])
        by_date.update({point["date"]: point for point in previous if point.get("date") and isinstance(point.get("value"), (int, float))})
    except (FileNotFoundError, json.JSONDecodeError):
        pass
    end = datetime.now(timezone.utc).date() - timedelta(days=1)
    cursor = end - timedelta(days=365)
    missing = []
    while cursor <= end:
        if cursor.weekday() < 5 and cursor.isoformat() not in by_date:
            missing.append(cursor.isoformat())
        cursor += timedelta(days=1)
    if missing:
        with ThreadPoolExecutor(max_workers=8) as pool:
            for future in as_completed([pool.submit(fetch_cboe_daily, item) for item in missing]):
                try:
                    point = future.result()
                    if point:
                        by_date[point["date"]] = point
                except Exception:
                    pass
    points = [by_date[key] for key in sorted(by_date)]
    return {
        "CBOE_PC": points_result(points),
        "CBOE_PC_MA20": points_result(moving_average_with_gap_reset(points, 20)),
    }

def main():
    data={"generatedAt":datetime.now(timezone.utc).isoformat(),"series":{}}
    def build_one(item):
        sid,(mode,scale)=item
        try:
            points=transform(fetch(sid),mode,scale)
            return sid,points_result(points)
        except Exception as exc:
            return sid,{"points":[],"error":f"{type(exc).__name__}: {exc}"}
    with ThreadPoolExecutor(max_workers=6) as pool:
        for future in as_completed([pool.submit(build_one,item) for item in SERIES.items()]):
            sid,result=future.result(); data["series"][sid]=result
    # Derived IDs let one source appear as a level and a YoY series without changing chart semantics.
    for source,derived in (("TOTALSA","TOTALSA_YOY"),("SPCS20RSA","SPCS20RSA_YOY")):
        try:
            points=transform(fetch(source),"yoy",1)
            data["series"][derived]=points_result(points)
        except Exception as exc: data["series"][derived]={"points":[],"error":str(exc)}
    for builder in (fetch_market_series, fetch_cftc_series, fetch_breadth_series, fetch_aaii_series, fetch_manheim_series, fetch_cboe_put_call):
        try:
            data["series"].update(builder())
        except Exception as exc:
            print(f"warning: {builder.__name__} failed: {type(exc).__name__}: {exc}")
    OUT.parent.mkdir(parents=True,exist_ok=True); OUT.write_text(json.dumps(data,separators=(",",":")),encoding="utf-8")
    print(f"wrote {OUT} ({OUT.stat().st_size/1024:.1f} KiB)")

if __name__ == "__main__": main()
